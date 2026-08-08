from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

from fastapi import FastAPI, APIRouter, Depends, HTTPException, Request, Query
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Literal
from datetime import datetime, timezone
import uuid
import io
import logging
import jwt
import bcrypt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ------------------------------------------------------------------ DB
mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

JWT_SECRET = os.environ["JWT_SECRET"]
JWT_ALGORITHM = "HS256"

app = FastAPI(title="EventGudang API")
api = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("eventgudang")


# ------------------------------------------------------------------ helpers
def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def new_id() -> str:
    return str(uuid.uuid4())


def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))
    except Exception:
        return False


def create_access_token(user_id: str, email: str, role: str) -> str:
    payload = {
        "sub": user_id,
        "email": email,
        "role": role,
        "exp": datetime.now(timezone.utc).timestamp() + 60 * 60 * 24 * 7,
        "type": "access",
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


async def get_current_user(request: Request) -> dict:
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else request.cookies.get("access_token")
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sesi berakhir, silakan masuk kembali")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "hashed_password": 0})
    if not user or not user.get("is_active", True):
        raise HTTPException(status_code=401, detail="Pengguna tidak ditemukan")
    return user


async def require_admin(user: dict = Depends(get_current_user)) -> dict:
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Akses khusus admin")
    return user


# ------------------------------------------------------------------ models
class LoginIn(BaseModel):
    email: EmailStr
    password: str


class UserOut(BaseModel):
    id: str
    email: str
    full_name: str
    role: str
    is_active: bool
    created_at: str


class UserCreate(BaseModel):
    email: EmailStr
    full_name: str
    password: str
    role: Literal["admin", "staff"] = "staff"


class WarehouseIn(BaseModel):
    name: str
    address: Optional[str] = ""


class BrandIn(BaseModel):
    name: str
    color: str = "#3B82F6"
    note: Optional[str] = ""


class CategoryIn(BaseModel):
    name: str


class ItemIn(BaseModel):
    name: str
    sku: str
    warehouse_id: str
    brand_id: Optional[str] = None
    category_id: Optional[str] = None
    total_qty: int = 0
    photo: Optional[str] = None
    notes: Optional[str] = ""


class ItemUpdate(BaseModel):
    name: Optional[str] = None
    sku: Optional[str] = None
    warehouse_id: Optional[str] = None
    brand_id: Optional[str] = None
    category_id: Optional[str] = None
    total_qty: Optional[int] = None
    photo: Optional[str] = None
    notes: Optional[str] = None


class SJItemLineIn(BaseModel):
    item_id: str
    qty: int


class SuratJalanIn(BaseModel):
    warehouse_id: str
    recipient_name: str
    event_name: str
    event_date: Optional[str] = ""
    notes: Optional[str] = ""
    items: List[SJItemLineIn]


class ReturnLine(BaseModel):
    item_id: str
    returned_qty: int = 0
    damaged_qty: int = 0
    lost_qty: int = 0
    used_qty: int = 0


class ReturnIn(BaseModel):
    lines: List[ReturnLine]


class PBLineIn(BaseModel):
    mode: Literal["existing", "new"]
    item_id: Optional[str] = None
    name: Optional[str] = None
    sku: Optional[str] = None
    brand_id: Optional[str] = None
    category_id: Optional[str] = None
    qty: int


class PenerimaanIn(BaseModel):
    warehouse_id: str
    supplier_name: str
    receipt_date: Optional[str] = ""
    notes: Optional[str] = ""
    items: List[PBLineIn]


# ------------------------------------------------------------------ auth routes
@api.post("/auth/login")
async def login(body: LoginIn):
    email = body.email.lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(body.password, user["hashed_password"]):
        raise HTTPException(status_code=401, detail="Email atau kata sandi salah")
    if not user.get("is_active", True):
        raise HTTPException(status_code=403, detail="Akun dinonaktifkan")
    token = create_access_token(user["id"], user["email"], user["role"])
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {
            "id": user["id"],
            "email": user["email"],
            "full_name": user["full_name"],
            "role": user["role"],
            "is_active": user.get("is_active", True),
            "created_at": user["created_at"],
        },
    }


@api.get("/auth/me", response_model=UserOut)
async def me(user: dict = Depends(get_current_user)):
    return user


# ------------------------------------------------------------------ users
@api.get("/users", response_model=List[UserOut])
async def list_users(_: dict = Depends(require_admin)):
    return await db.users.find({}, {"_id": 0, "hashed_password": 0}).sort("created_at", -1).to_list(500)


@api.post("/users", response_model=UserOut)
async def create_user(body: UserCreate, _: dict = Depends(require_admin)):
    email = body.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    doc = {
        "id": new_id(),
        "email": email,
        "full_name": body.full_name,
        "hashed_password": hash_password(body.password),
        "role": body.role,
        "is_active": True,
        "created_at": now_iso(),
    }
    await db.users.insert_one(doc)
    doc.pop("_id", None)
    doc.pop("hashed_password", None)
    return doc


# ------------------------------------------------------------------ warehouses
@api.get("/warehouses")
async def list_warehouses(_: dict = Depends(get_current_user)):
    return await db.warehouses.find({}, {"_id": 0}).sort("name", 1).to_list(500)


@api.post("/warehouses")
async def create_warehouse(body: WarehouseIn, _: dict = Depends(require_admin)):
    if await db.warehouses.find_one({"name": body.name}):
        raise HTTPException(status_code=400, detail="Nama gudang sudah ada")
    doc = {"id": new_id(), "name": body.name, "address": body.address or "", "created_at": now_iso()}
    await db.warehouses.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api.patch("/warehouses/{wid}")
async def update_warehouse(wid: str, body: WarehouseIn, _: dict = Depends(require_admin)):
    res = await db.warehouses.update_one({"id": wid}, {"$set": {"name": body.name, "address": body.address or ""}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Gudang tidak ditemukan")
    return await db.warehouses.find_one({"id": wid}, {"_id": 0})


@api.delete("/warehouses/{wid}")
async def delete_warehouse(wid: str, _: dict = Depends(require_admin)):
    if await db.items.find_one({"warehouse_id": wid}):
        raise HTTPException(status_code=400, detail="Gudang masih memiliki barang")
    await db.warehouses.delete_one({"id": wid})
    return {"ok": True}


# ------------------------------------------------------------------ brands
@api.get("/brands")
async def list_brands(_: dict = Depends(get_current_user)):
    return await db.brands.find({}, {"_id": 0}).sort("name", 1).to_list(500)


@api.post("/brands")
async def create_brand(body: BrandIn, _: dict = Depends(require_admin)):
    if await db.brands.find_one({"name": body.name}):
        raise HTTPException(status_code=400, detail="Nama brand sudah ada")
    doc = {"id": new_id(), "name": body.name, "color": body.color, "note": body.note or "", "created_at": now_iso()}
    await db.brands.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api.patch("/brands/{bid}")
async def update_brand(bid: str, body: BrandIn, _: dict = Depends(require_admin)):
    res = await db.brands.update_one(
        {"id": bid}, {"$set": {"name": body.name, "color": body.color, "note": body.note or ""}}
    )
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Brand tidak ditemukan")
    return await db.brands.find_one({"id": bid}, {"_id": 0})


@api.delete("/brands/{bid}")
async def delete_brand(bid: str, _: dict = Depends(require_admin)):
    if await db.items.find_one({"brand_id": bid}):
        raise HTTPException(status_code=400, detail="Brand masih dipakai oleh barang")
    await db.brands.delete_one({"id": bid})
    return {"ok": True}


# ------------------------------------------------------------------ categories
@api.get("/categories")
async def list_categories(_: dict = Depends(get_current_user)):
    return await db.categories.find({}, {"_id": 0}).sort("name", 1).to_list(500)


@api.post("/categories")
async def create_category(body: CategoryIn, _: dict = Depends(require_admin)):
    if await db.categories.find_one({"name": body.name}):
        raise HTTPException(status_code=400, detail="Jenis sudah ada")
    doc = {"id": new_id(), "name": body.name, "created_at": now_iso()}
    await db.categories.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api.patch("/categories/{cid}")
async def update_category(cid: str, body: CategoryIn, _: dict = Depends(require_admin)):
    res = await db.categories.update_one({"id": cid}, {"$set": {"name": body.name}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Jenis tidak ditemukan")
    return await db.categories.find_one({"id": cid}, {"_id": 0})


@api.delete("/categories/{cid}")
async def delete_category(cid: str, _: dict = Depends(require_admin)):
    if await db.items.find_one({"category_id": cid}):
        raise HTTPException(status_code=400, detail="Jenis masih dipakai oleh barang")
    await db.categories.delete_one({"id": cid})
    return {"ok": True}


# ------------------------------------------------------------------ items
async def build_ref_maps():
    whs = await db.warehouses.find({}, {"_id": 0, "id": 1, "name": 1}).to_list(2000)
    brs = await db.brands.find({}, {"_id": 0, "id": 1, "name": 1, "color": 1}).to_list(2000)
    cats = await db.categories.find({}, {"_id": 0, "id": 1, "name": 1}).to_list(2000)
    return ({w["id"]: w for w in whs}, {b["id"]: b for b in brs}, {c["id"]: c for c in cats})


def enrich_with_maps(item, wmap, bmap, cmap):
    w = wmap.get(item.get("warehouse_id"))
    b = bmap.get(item.get("brand_id"))
    c = cmap.get(item.get("category_id"))
    item["warehouse_name"] = w["name"] if w else None
    item["brand_name"] = b["name"] if b else None
    item["brand_color"] = b["color"] if b else None
    item["category_name"] = c["name"] if c else None
    return item


async def enrich_item(item: dict) -> dict:
    wh = await db.warehouses.find_one({"id": item.get("warehouse_id")}, {"_id": 0, "name": 1})
    br = await db.brands.find_one({"id": item.get("brand_id")}, {"_id": 0, "name": 1, "color": 1})
    cat = await db.categories.find_one({"id": item.get("category_id")}, {"_id": 0, "name": 1})
    item["warehouse_name"] = wh["name"] if wh else None
    item["brand_name"] = br["name"] if br else None
    item["brand_color"] = br["color"] if br else None
    item["category_name"] = cat["name"] if cat else None
    return item


@api.get("/items")
async def list_items(
    warehouse_id: Optional[str] = None,
    brand_id: Optional[str] = None,
    category_id: Optional[str] = None,
    _: dict = Depends(get_current_user),
):
    q = {}
    if warehouse_id:
        q["warehouse_id"] = warehouse_id
    if brand_id:
        q["brand_id"] = brand_id
    if category_id:
        q["category_id"] = category_id
    items = await db.items.find(q, {"_id": 0}).sort("created_at", -1).to_list(1000)
    wmap, bmap, cmap = await build_ref_maps()
    for it in items:
        enrich_with_maps(it, wmap, bmap, cmap)
    return items


@api.get("/items/{item_id}")
async def get_item(item_id: str, _: dict = Depends(get_current_user)):
    item = await db.items.find_one({"id": item_id}, {"_id": 0})
    if not item:
        raise HTTPException(status_code=404, detail="Barang tidak ditemukan")
    return await enrich_item(item)


@api.post("/items")
async def create_item(body: ItemIn, _: dict = Depends(require_admin)):
    if await db.items.find_one({"sku": body.sku}):
        raise HTTPException(status_code=400, detail="SKU sudah ada")
    doc = {
        "id": new_id(),
        "name": body.name,
        "sku": body.sku,
        "warehouse_id": body.warehouse_id,
        "brand_id": body.brand_id,
        "category_id": body.category_id,
        "total_qty": body.total_qty,
        "available_qty": body.total_qty,
        "damaged_qty": 0,
        "lost_qty": 0,
        "used_qty": 0,
        "photo": body.photo,
        "notes": body.notes or "",
        "created_at": now_iso(),
    }
    await db.items.insert_one(doc)
    doc.pop("_id", None)
    return await enrich_item(doc)


@api.patch("/items/{item_id}")
async def update_item(item_id: str, body: ItemUpdate, _: dict = Depends(require_admin)):
    item = await db.items.find_one({"id": item_id})
    if not item:
        raise HTTPException(status_code=404, detail="Barang tidak ditemukan")
    updates = body.model_dump(exclude_none=True)
    if "sku" in updates and updates["sku"] != item["sku"]:
        if await db.items.find_one({"sku": updates["sku"]}):
            raise HTTPException(status_code=400, detail="SKU sudah ada")
    if "total_qty" in updates:
        # adjust available by the delta of total
        delta = updates["total_qty"] - item["total_qty"]
        updates["available_qty"] = max(0, item["available_qty"] + delta)
    await db.items.update_one({"id": item_id}, {"$set": updates})
    doc = await db.items.find_one({"id": item_id}, {"_id": 0})
    return await enrich_item(doc)


@api.delete("/items/{item_id}")
async def delete_item(item_id: str, _: dict = Depends(require_admin)):
    await db.items.delete_one({"id": item_id})
    return {"ok": True}


# ------------------------------------------------------------------ surat jalan
async def gen_sj_number() -> str:
    prefix = "SJ-" + datetime.now(timezone.utc).strftime("%Y%m") + "-"
    count = await db.surat_jalan.count_documents({"sj_number": {"$regex": f"^{prefix}"}})
    return f"{prefix}{count + 1:04d}"


@api.get("/surat-jalan")
async def list_sj(status_filter: Optional[str] = Query(None), _: dict = Depends(get_current_user)):
    q = {}
    if status_filter == "aktif":
        q["status"] = {"$in": ["out", "partial"]}
    elif status_filter == "selesai":
        q["status"] = "returned"
    return await db.surat_jalan.find(q, {"_id": 0}).sort("created_at", -1).to_list(1000)


@api.get("/surat-jalan/by-token/{qr_token}")
async def sj_by_token(qr_token: str, _: dict = Depends(get_current_user)):
    sj = await db.surat_jalan.find_one({"qr_token": qr_token}, {"_id": 0})
    if not sj:
        raise HTTPException(status_code=404, detail="Surat jalan tidak ditemukan")
    return sj


@api.get("/surat-jalan/{sj_id}")
async def get_sj(sj_id: str, _: dict = Depends(get_current_user)):
    sj = await db.surat_jalan.find_one({"id": sj_id}, {"_id": 0})
    if not sj:
        raise HTTPException(status_code=404, detail="Surat jalan tidak ditemukan")
    return sj


@api.post("/surat-jalan")
async def create_sj(body: SuratJalanIn, user: dict = Depends(get_current_user)):
    wh = await db.warehouses.find_one({"id": body.warehouse_id}, {"_id": 0})
    if not wh:
        raise HTTPException(status_code=404, detail="Gudang tidak ditemukan")
    if not body.items:
        raise HTTPException(status_code=400, detail="Minimal satu barang")
    lines = []
    for line in body.items:
        item = await db.items.find_one({"id": line.item_id})
        if not item:
            raise HTTPException(status_code=404, detail="Barang tidak ditemukan")
        if line.qty <= 0:
            continue
        if item["available_qty"] < line.qty:
            raise HTTPException(status_code=400, detail=f"Stok {item['name']} tidak cukup")
        lines.append(
            {
                "item_id": item["id"],
                "item_name": item["name"],
                "sku": item["sku"],
                "qty": line.qty,
                "returned_qty": 0,
                "damaged_qty": 0,
                "lost_qty": 0,
                "used_qty": 0,
            }
        )
    if not lines:
        raise HTTPException(status_code=400, detail="Minimal satu barang dengan qty > 0")
    for line in lines:
        await db.items.update_one({"id": line["item_id"]}, {"$inc": {"available_qty": -line["qty"]}})
    doc = {
        "id": new_id(),
        "sj_number": await gen_sj_number(),
        "qr_token": new_id(),
        "warehouse_id": wh["id"],
        "warehouse_name": wh["name"],
        "recipient_name": body.recipient_name,
        "event_name": body.event_name,
        "event_date": body.event_date or "",
        "notes": body.notes or "",
        "items": lines,
        "status": "out",
        "created_by": user["id"],
        "created_by_name": user["full_name"],
        "created_at": now_iso(),
        "returned_at": None,
    }
    await db.surat_jalan.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api.post("/surat-jalan/{sj_id}/return")
async def return_sj(sj_id: str, body: ReturnIn, _: dict = Depends(get_current_user)):
    sj = await db.surat_jalan.find_one({"id": sj_id})
    if not sj:
        raise HTTPException(status_code=404, detail="Surat jalan tidak ditemukan")
    if sj["status"] == "returned":
        raise HTTPException(status_code=400, detail="Surat jalan sudah selesai")
    line_map = {l["item_id"]: l for l in sj["items"]}
    for rl in body.lines:
        line = line_map.get(rl.item_id)
        if not line:
            continue
        outstanding = line["qty"] - (line["returned_qty"] + line["damaged_qty"] + line["lost_qty"] + line["used_qty"])
        total_now = rl.returned_qty + rl.damaged_qty + rl.lost_qty + rl.used_qty
        if total_now <= 0:
            continue
        if total_now > outstanding:
            raise HTTPException(status_code=400, detail=f"Jumlah pengembalian {line['item_name']} melebihi sisa")
        # apply stock effects
        inc = {}
        if rl.returned_qty:
            inc["available_qty"] = inc.get("available_qty", 0) + rl.returned_qty
        if rl.damaged_qty:
            inc["damaged_qty"] = inc.get("damaged_qty", 0) + rl.damaged_qty
        if rl.lost_qty:
            inc["lost_qty"] = inc.get("lost_qty", 0) + rl.lost_qty
            inc["total_qty"] = inc.get("total_qty", 0) - rl.lost_qty
        if rl.used_qty:
            inc["used_qty"] = inc.get("used_qty", 0) + rl.used_qty
            inc["total_qty"] = inc.get("total_qty", 0) - rl.used_qty
        if inc:
            await db.items.update_one({"id": rl.item_id}, {"$inc": inc})
        line["returned_qty"] += rl.returned_qty
        line["damaged_qty"] += rl.damaged_qty
        line["lost_qty"] += rl.lost_qty
        line["used_qty"] += rl.used_qty
    # recompute status
    fully_done = all(
        (l["returned_qty"] + l["damaged_qty"] + l["lost_qty"] + l["used_qty"]) >= l["qty"] for l in sj["items"]
    )
    any_done = any(
        (l["returned_qty"] + l["damaged_qty"] + l["lost_qty"] + l["used_qty"]) > 0 for l in sj["items"]
    )
    status = "returned" if fully_done else ("partial" if any_done else "out")
    update = {"items": sj["items"], "status": status}
    if status == "returned":
        update["returned_at"] = now_iso()
    await db.surat_jalan.update_one({"id": sj_id}, {"$set": update})
    return await db.surat_jalan.find_one({"id": sj_id}, {"_id": 0})


# ------------------------------------------------------------------ penerimaan
async def gen_pbk_number() -> str:
    prefix = "PBK-" + datetime.now(timezone.utc).strftime("%Y%m") + "-"
    count = await db.penerimaan.count_documents({"receipt_number": {"$regex": f"^{prefix}"}})
    return f"{prefix}{count + 1:04d}"


@api.get("/penerimaan")
async def list_penerimaan(_: dict = Depends(get_current_user)):
    return await db.penerimaan.find({}, {"_id": 0}).sort("created_at", -1).to_list(1000)


@api.get("/penerimaan/{pid}")
async def get_penerimaan(pid: str, _: dict = Depends(get_current_user)):
    p = await db.penerimaan.find_one({"id": pid}, {"_id": 0})
    if not p:
        raise HTTPException(status_code=404, detail="Penerimaan tidak ditemukan")
    return p


@api.post("/penerimaan")
async def create_penerimaan(body: PenerimaanIn, user: dict = Depends(get_current_user)):
    wh = await db.warehouses.find_one({"id": body.warehouse_id}, {"_id": 0})
    if not wh:
        raise HTTPException(status_code=404, detail="Gudang tidak ditemukan")
    if not body.items:
        raise HTTPException(status_code=400, detail="Minimal satu barang")
    lines = []
    for line in body.items:
        if line.qty <= 0:
            continue
        if line.mode == "existing":
            item = await db.items.find_one({"id": line.item_id})
            if not item:
                raise HTTPException(status_code=404, detail="Barang tidak ditemukan")
            await db.items.update_one(
                {"id": item["id"]}, {"$inc": {"total_qty": line.qty, "available_qty": line.qty}}
            )
            lines.append(
                {
                    "mode": "existing",
                    "item_id": item["id"],
                    "item_name": item["name"],
                    "sku": item["sku"],
                    "qty": line.qty,
                    "is_new": False,
                }
            )
        else:  # new
            if not line.name or not line.sku:
                raise HTTPException(status_code=400, detail="Nama & SKU barang baru wajib diisi")
            if await db.items.find_one({"sku": line.sku}):
                raise HTTPException(status_code=400, detail=f"SKU {line.sku} sudah ada")
            new_item = {
                "id": new_id(),
                "name": line.name,
                "sku": line.sku,
                "warehouse_id": wh["id"],
                "brand_id": line.brand_id,
                "category_id": line.category_id,
                "total_qty": line.qty,
                "available_qty": line.qty,
                "damaged_qty": 0,
                "lost_qty": 0,
                "used_qty": 0,
                "photo": None,
                "notes": "",
                "created_at": now_iso(),
            }
            await db.items.insert_one(new_item)
            lines.append(
                {
                    "mode": "new",
                    "item_id": new_item["id"],
                    "item_name": new_item["name"],
                    "sku": new_item["sku"],
                    "qty": line.qty,
                    "is_new": True,
                }
            )
    if not lines:
        raise HTTPException(status_code=400, detail="Minimal satu barang dengan qty > 0")
    doc = {
        "id": new_id(),
        "receipt_number": await gen_pbk_number(),
        "warehouse_id": wh["id"],
        "warehouse_name": wh["name"],
        "supplier_name": body.supplier_name,
        "receipt_date": body.receipt_date or "",
        "notes": body.notes or "",
        "items": lines,
        "created_by": user["id"],
        "created_by_name": user["full_name"],
        "created_at": now_iso(),
    }
    await db.penerimaan.insert_one(doc)
    doc.pop("_id", None)
    return doc


# ------------------------------------------------------------------ dashboard
@api.get("/dashboard")
async def dashboard(_: dict = Depends(get_current_user)):
    items = await db.items.find(
        {},
        {"_id": 0, "total_qty": 1, "available_qty": 1, "damaged_qty": 1, "lost_qty": 1, "used_qty": 1},
    ).to_list(5000)
    available = sum(i.get("available_qty", 0) for i in items)
    total_units = sum(i.get("total_qty", 0) for i in items)
    damaged = sum(i.get("damaged_qty", 0) for i in items)
    lost = sum(i.get("lost_qty", 0) for i in items)
    used = sum(i.get("used_qty", 0) for i in items)
    keluar = sum(max(0, i.get("total_qty", 0) - i.get("available_qty", 0)) for i in items)
    wh_count = await db.warehouses.count_documents({})
    sj_active = await db.surat_jalan.count_documents({"status": {"$in": ["out", "partial"]}})
    sj_done = await db.surat_jalan.count_documents({"status": "returned"})
    recent = await db.surat_jalan.find({}, {"_id": 0}).sort("created_at", -1).to_list(5)
    return {
        "jenis_barang": len(items),
        "tersedia": available,
        "keluar": keluar,
        "total_unit": total_units,
        "gudang": wh_count,
        "rusak": damaged,
        "hilang": lost,
        "terpakai": used,
        "sj_aktif": sj_active,
        "sj_selesai": sj_done,
        "recent_surat_jalan": recent,
    }


# ------------------------------------------------------------------ excel export
@api.get("/export/inventory.xlsx")
async def export_inventory(_: dict = Depends(require_admin)):
    items = await db.items.find({}, {"_id": 0}).to_list(5000)
    wmap, bmap, cmap = await build_ref_maps()
    for it in items:
        enrich_with_maps(it, wmap, bmap, cmap)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventori"
    headers = ["SKU", "Nama", "Brand", "Jenis", "Gudang", "Total", "Tersedia", "Keluar", "Rusak", "Hilang", "Terpakai", "Catatan"]
    ws.append(headers)
    header_fill = PatternFill(start_color="002FA7", end_color="002FA7", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for it in items:
        keluar = max(0, it.get("total_qty", 0) - it.get("available_qty", 0))
        ws.append([
            it.get("sku", ""), it.get("name", ""), it.get("brand_name") or "", it.get("category_name") or "",
            it.get("warehouse_name") or "", it.get("total_qty", 0), it.get("available_qty", 0), keluar,
            it.get("damaged_qty", 0), it.get("lost_qty", 0), it.get("used_qty", 0), it.get("notes", ""),
        ])
    widths = [18, 30, 16, 14, 18, 8, 10, 8, 8, 8, 10, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"inventori-{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


@api.get("/")
async def root():
    return {"message": "EventGudang API"}


app.include_router(api)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get("CORS_ORIGINS", "*").split(","),
    allow_methods=["*"],
    allow_headers=["*"],
)


# ------------------------------------------------------------------ seeding
SEED_BRANDS = [
    ("JBL", "#E63946"), ("Yamaha", "#457B9D"), ("Shure", "#2A9D8F"),
    ("Chauvet", "#E76F51"), ("Global Truss", "#6A4C93"), ("Custom Merch", "#F4A261"),
]
SEED_CATEGORIES = ["Tools", "Merchandise", "Produk"]
SEED_WAREHOUSES = [("Gudang Pusat", "Jl. Industri No. 1, Jakarta"), ("Gudang Bandung", "Jl. Soekarno Hatta No. 88, Bandung")]


async def seed():
    # users
    for email_key, pw_key, name, role in [
        ("ADMIN_EMAIL", "ADMIN_PASSWORD", "Administrator", "admin"),
        ("STAFF_EMAIL", "STAFF_PASSWORD", "Staff Gudang", "staff"),
    ]:
        email = os.environ.get(email_key, "").lower()
        pw = os.environ.get(pw_key, "")
        if not email or not pw:
            continue
        existing = await db.users.find_one({"email": email})
        if not existing:
            await db.users.insert_one({
                "id": new_id(), "email": email, "full_name": name,
                "hashed_password": hash_password(pw), "role": role,
                "is_active": True, "created_at": now_iso(),
            })
        elif not verify_password(pw, existing["hashed_password"]):
            await db.users.update_one({"email": email}, {"$set": {"hashed_password": hash_password(pw)}})

    if await db.items.count_documents({}) > 0:
        return  # data already seeded

    wh_ids = {}
    for name, addr in SEED_WAREHOUSES:
        wid = new_id()
        wh_ids[name] = wid
        await db.warehouses.insert_one({"id": wid, "name": name, "address": addr, "created_at": now_iso()})

    br_ids = {}
    for name, color in SEED_BRANDS:
        bid = new_id()
        br_ids[name] = bid
        await db.brands.insert_one({"id": bid, "name": name, "color": color, "note": "", "created_at": now_iso()})

    cat_ids = {}
    for name in SEED_CATEGORIES:
        cid = new_id()
        cat_ids[name] = cid
        await db.categories.insert_one({"id": cid, "name": name, "created_at": now_iso()})

    seed_items = [
        ("Speaker EON 715", "SPK-EON715", "Gudang Pusat", "JBL", "Tools", 12),
        ("Mixer MG12XU", "MIX-MG12XU", "Gudang Pusat", "Yamaha", "Tools", 4),
        ("Wireless Mic BLX", "MIC-BLX24", "Gudang Pusat", "Shure", "Tools", 20),
        ("Moving Head Beam 230", "LGT-BEAM230", "Gudang Pusat", "Chauvet", "Tools", 8),
        ("LED Par 64 RGBW", "LGT-PAR64", "Gudang Bandung", "Chauvet", "Tools", 24),
        ("Truss Frame 3m", "TRS-FRAME3M", "Gudang Bandung", "Global Truss", "Tools", 30),
        ("Truss Base Plate", "TRS-BASE", "Gudang Bandung", "Global Truss", "Tools", 16),
        ("Subwoofer SRX828", "SPK-SRX828", "Gudang Bandung", "JBL", "Tools", 6),
        ("T-Shirt Event XL", "MRC-TSHIRT-XL", "Gudang Pusat", "Custom Merch", "Merchandise", 200),
        ("Tote Bag Canvas", "MRC-TOTE", "Gudang Pusat", "Custom Merch", "Merchandise", 150),
        ("Lanyard ID Card", "MRC-LANYARD", "Gudang Bandung", "Custom Merch", "Produk", 500),
    ]
    for name, sku, wh, br, cat, qty in seed_items:
        await db.items.insert_one({
            "id": new_id(), "name": name, "sku": sku, "warehouse_id": wh_ids[wh],
            "brand_id": br_ids[br], "category_id": cat_ids[cat],
            "total_qty": qty, "available_qty": qty, "damaged_qty": 0, "lost_qty": 0, "used_qty": 0,
            "photo": None, "notes": "", "created_at": now_iso(),
        })
    logger.info("Seed data inserted")


@app.on_event("startup")
async def on_startup():
    await db.users.create_index("email", unique=True)
    await db.warehouses.create_index("name", unique=True)
    await db.brands.create_index("name", unique=True)
    await db.categories.create_index("name", unique=True)
    await db.items.create_index("sku", unique=True)
    await db.surat_jalan.create_index("qr_token")
    await seed()


@app.on_event("shutdown")
async def shutdown():
    client.close()
